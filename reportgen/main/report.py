from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from io import BytesIO
from .logger import logging
class Report:
    def __init__(self):
        self.doc = Document()
    def add_title_page(self,title):
        logging.info("Adding title page")
        para = self.doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run(title)
        run.font.size = Pt(24)
        self.doc.add_page_break()
        logging.info("Adding title page completed")
    def add_content(self,heading,content):
        logging.info(f"Adding content of {heading}")
        self.doc.add_heading(heading,level=1)
        self.doc.add_paragraph(content)

    def return_doc(self):
        return self.doc
    
    def convert_bytes(self):
        logging.info("Document to bytes conversion started")
        f=BytesIO()
        self.doc.save(f)
        f.seek(0)
        logging.info("Bytes file returning")
        return f

